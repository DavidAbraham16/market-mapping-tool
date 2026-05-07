import os
import io
import math
import json
from flask import Flask, request, jsonify, send_file, render_template
import anthropic
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

app = Flask(__name__)
client = anthropic.Anthropic(api_key=os.environ.get("ANTHROPIC_API_KEY"))

NAVY = "FF1F3864"
BLUE = "FF2E75B6"
WHITE = "FFFFFFFF"
GREEN_BG = "FFE2EFDA"
GREEN_FG = "FF375623"
RED_BG = "FFFCE4D6"
RED_FG = "FF9C0006"

OFFICES = [
    ("Chicago (Schaumburg, IL)", 42.03, -88.08),
    ("Massachusetts (Needham, MA)", 42.28, -71.23),
    ("North Carolina (Chapel Hill)", 35.91, -79.06),
    ("Los Angeles (Long Beach, CA)", 33.77, -118.19),
    ("San Francisco (San Ramon, CA)", 37.78, -121.98),
]

CITY_COORDS = {
    "schaumburg": (42.03, -88.08), "chicago": (41.88, -87.63),
    "needham": (42.28, -71.23), "boston": (42.36, -71.06),
    "westford": (42.58, -71.43), "boston": (42.36, -71.06),
    "chapel hill": (35.91, -79.06), "raleigh": (35.78, -78.64),
    "charlotte": (35.23, -80.84),
    "long beach": (33.77, -118.19), "los angeles": (34.05, -118.24),
    "san diego": (32.72, -117.16),
    "san ramon": (37.78, -121.98), "san francisco": (37.77, -122.42),
    "denver": (39.74, -104.99), "seattle": (47.61, -122.33),
    "indianapolis": (39.77, -86.16), "atlanta": (33.75, -84.39),
    "cincinnati": (39.10, -84.51), "columbus": (39.96, -82.99),
    "dallas": (32.78, -96.80), "houston": (29.76, -95.37),
    "austin": (30.27, -97.74), "nashville": (36.16, -86.78),
    "franklin": (35.93, -86.87), "boca raton": (26.37, -80.13),
    "miami": (25.76, -80.19), "tampa": (27.95, -82.46),
    "new york": (40.71, -74.01), "albany": (42.65, -73.76),
    "philadelphia": (39.95, -75.17), "pittsburgh": (40.44, -79.99),
    "hanover": (39.81, -76.98), "waterford": (39.12, -77.55),
    "minneapolis": (44.98, -93.27), "kansas city": (39.10, -94.58),
    "st. louis": (38.63, -90.20), "phoenix": (33.45, -112.07),
    "portland": (45.51, -122.68), "salt lake city": (40.76, -111.89),
    "mason": (39.36, -84.31), "west chester": (39.56, -84.39),
}

def haversine(lat1, lon1, lat2, lon2):
    R = 6371
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dp, dl = math.radians(lat2 - lat1), math.radians(lon2 - lon1)
    a = math.sin(dp/2)**2 + math.cos(p1) * math.cos(p2) * math.sin(dl/2)**2
    return 2 * R * math.atan2(math.sqrt(a), math.sqrt(1 - a))

def get_nearest_office(location):
    if not location:
        return "North Carolina (Chapel Hill)", "Remote"
    loc = location.lower()
    if any(x in loc for x in ["uk", "united kingdom", "england", "germany", "france", "australia", "canada", "ireland"]):
        return "International", "Remote"
    coords = None
    for city, c in CITY_COORDS.items():
        if city in loc:
            coords = c
            break
    if not coords:
        return "North Carolina (Chapel Hill)", "Remote"
    nearest, min_d = None, float("inf")
    for name, lat, lon in OFFICES:
        d = haversine(coords[0], coords[1], lat, lon)
        if d < min_d:
            min_d = d
            nearest = name
    return nearest, ("Hybrid" if min_d <= 100 else "Remote")

def thin_border():
    s = Side(style="thin", color="FFD9D9D9")
    return Border(left=s, right=s, top=s, bottom=s)

def parse_candidate(company_name, profile_text):
    prompt = f"""Parse this LinkedIn profile for a warehouse automation market mapping report.
Return ONLY valid JSON, no markdown, no extra text.

{{
  "name": "Full name",
  "years": "Total years experience as string e.g. ~8 years",
  "currentRole": "Current job title at Company",
  "previousRole": "Most recent previous job title at Company",
  "overview": "Exactly 2 bullet points. Each starts with •. Max 15 words each. Focused on automation/robotics exposure. Separated by \\n",
  "location": "City, State, USA or City, Country if international"
}}

Company context: {company_name}
Profile:
{profile_text[:4000]}"""

    message = client.messages.create(
        model="claude-opus-4-5",
        max_tokens=800,
        messages=[{"role": "user", "content": prompt}]
    )
    text = message.content[0].text.strip()
    text = text.replace("```json", "").replace("```", "").strip()
    return json.loads(text)

def generate_company_overview(company_name, website=""):
    prompt = f"""Write a single sentence (max 25 words) describing what {company_name} does, 
focusing on their warehouse automation, robotics, or supply chain technology offering.
{'Their website is ' + website if website else ''}
Return ONLY the sentence, no quotes, no extra text."""
    message = client.messages.create(
        model="claude-opus-4-5",
        max_tokens=100,
        messages=[{"role": "user", "content": prompt}]
    )
    return message.content[0].text.strip().strip('"').strip("'")

def build_excel(market_name, companies):
    wb = Workbook()
    ws = wb.active
    ws.title = "Market Mapping"

    col_widths = [25, 18, 35, 35, 55, 25, 30]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    r = 1
    # Main header
    ws.row_dimensions[r].height = 30
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    c = ws.cell(row=r, column=1, value=market_name)
    c.font = Font(name="Arial", bold=True, size=14, color=WHITE)
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(horizontal="left", vertical="center")
    r += 1

    for company in companies:
        # Spacer
        ws.row_dimensions[r].height = 6
        r += 1

        # Company banner
        banner_text = f"{company['name']}  |  {company['overview']}"
        ws.row_dimensions[r].height = 30
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
        c = ws.cell(row=r, column=1, value=banner_text)
        c.font = Font(name="Arial", size=10, color=WHITE)
        c.fill = PatternFill("solid", fgColor=BLUE)
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        r += 1

        # Column headers
        headers = ["Name", "Years of Experience", "Current Role", "Previous Role",
                   "Career Overview", "Current Location", "Nearest WiseTech Office"]
        ws.row_dimensions[r].height = 20
        for i, h in enumerate(headers, 1):
            c = ws.cell(row=r, column=i, value=h)
            c.font = Font(name="Arial", bold=True, size=11, color=WHITE)
            c.fill = PatternFill("solid", fgColor=NAVY)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        r += 1

        # Sort candidates by years descending
        def extract_years(cand):
            import re
            m = re.search(r"(\d+)", cand.get("years", "0"))
            return int(m.group(1)) if m else 0

        candidates = sorted(company["candidates"], key=extract_years, reverse=True)

        for cand in candidates:
            office_name, office_status = get_nearest_office(cand.get("location", ""))
            ws.row_dimensions[r].height = 55

            vals = [
                cand.get("name", ""),
                cand.get("years", ""),
                cand.get("currentRole", ""),
                cand.get("previousRole", ""),
                cand.get("overview", ""),
                cand.get("location", ""),
            ]

            for i, val in enumerate(vals, 1):
                c = ws.cell(row=r, column=i, value=val)
                c.font = Font(name="Arial", size=10, color="FF000000")
                c.fill = PatternFill(fill_type=None)
                c.alignment = Alignment(vertical="top", wrap_text=True)
                c.border = thin_border()

            # Office cell
            oc = ws.cell(row=r, column=7, value=f"{office_name}\n{office_status}")
            if office_status == "Hybrid":
                oc.font = Font(name="Arial", bold=True, size=10, color=GREEN_FG)
                oc.fill = PatternFill("solid", fgColor=GREEN_BG)
            else:
                oc.font = Font(name="Arial", bold=True, size=10, color=RED_FG)
                oc.fill = PatternFill("solid", fgColor=RED_BG)
            oc.alignment = Alignment(vertical="top", wrap_text=True)
            oc.border = thin_border()
            r += 1

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

@app.route("/")
def index():
    return render_template("index.html")

@app.route("/build", methods=["POST"])
def build():
    data = request.json
    market_name = data.get("marketName", "Market Mapping")
    companies_input = data.get("companies", [])

    results = []
    for co in companies_input:
        name = co.get("name", "").strip()
        website = co.get("website", "").strip()
        overview = co.get("overview", "").strip()
        profiles = co.get("profiles", [])

        if not name:
            continue

        if not overview:
            overview = generate_company_overview(name, website)

        candidates = []
        for profile in profiles:
            if len(profile.strip()) < 50:
                continue
            try:
                parsed = parse_candidate(name, profile)
                candidates.append(parsed)
            except Exception as e:
                print(f"Error parsing profile: {e}")
                continue

        if candidates:
            results.append({"name": name, "overview": overview, "candidates": candidates})

    if not results:
        return jsonify({"error": "No valid companies or profiles found."}), 400

    excel_file = build_excel(market_name, results)
    filename = market_name.replace(" ", "_").replace("/", "-") + ".xlsx"
    return send_file(
        excel_file,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=8080, debug=False)
